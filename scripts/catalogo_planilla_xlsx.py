"""catalogo_planilla_xlsx.py — vuelca `docs/catalogo_series.csv` en la hoja «Catálogo» del xlsx.

Reescribe solo esa hoja del libro `docs/catalogo_series.xlsx` (las otras cuatro —«Cómo usar»,
«Resumen», «Fuentes», «Prioridades»— se conservan) y estira a la cantidad de filas nueva:
el autofiltro, las listas desplegables de las columnas amarillas y los rangos de las fórmulas
del resumen, que están escritos con la última fila fija.

    python scripts/catalogo_planilla_xlsx.py
"""

from __future__ import annotations

import csv
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parent.parent
CSV = RAIZ / "docs" / "catalogo_series.csv"
XLSX = RAIZ / "docs" / "catalogo_series.xlsx"
NUMERICAS = {"Obs"}


def main() -> None:
    with CSV.open(encoding="utf-8", newline="") as fh:
        filas = list(csv.DictReader(fh))
    cab = list(filas[0])
    ultima = len(filas) + 1

    wb = openpyxl.load_workbook(XLSX)
    hoja = wb["Catálogo"]
    previa = hoja.max_row

    # La cabecera de la hoja se reescribe siempre. Si cambió el juego de columnas —agregar
    # «Grupo» corrió todo un lugar— hay que remapear las fórmulas de las otras hojas, que
    # referencian columnas por letra: Estado pasó de M a N, Revisado de R a S. Sin esto la
    # planilla queda silenciosamente desalineada y los conteos del resumen cuentan otra columna.
    cab_previa = [str(c.value) if c.value is not None else ""
                  for c in next(hoja.iter_rows(min_row=1, max_row=1))]
    letras = {get_column_letter(i + 1): get_column_letter(cab.index(n) + 1)
              for i, n in enumerate(cab_previa) if n in cab}
    if cab_previa[:len(cab)] != cab:
        movidas = {a: b for a, b in letras.items() if a != b}
        for h in wb.worksheets:
            if h is hoja:
                continue
            for fila in h.iter_rows():
                for c in fila:
                    if isinstance(c.value, str) and "Catálogo!" in c.value:
                        # Se remapea el rango COMPLETO: `Catálogo!$M$2:$M$1091` lleva la letra
                        # dos veces y solo la primera trae el prefijo de la hoja.
                        c.value = re.sub(
                            r"(Catálogo!\$)([A-Z]+)(\$\d+:\$)([A-Z]+)(\$)",
                            lambda m: (m.group(1) + letras.get(m.group(2), m.group(2))
                                       + m.group(3) + letras.get(m.group(4), m.group(4))
                                       + m.group(5)),
                            c.value)
        print(f"cabecera nueva: {len(cab_previa)} -> {len(cab)} columnas · "
              f"columnas movidas {movidas or '—'}")
    for j, col in enumerate(cab, start=1):
        hoja.cell(row=1, column=j, value=col)

    estilo = {c: hoja.cell(row=2, column=c)._style for c in range(1, len(cab) + 1)}
    validaciones = [(str(next(iter(dv.sqref.ranges))), dv.formula1)
                    for dv in list(hoja.data_validations.dataValidation)]

    hoja.delete_rows(2, max(previa - 1, 0))
    for i, f in enumerate(filas, start=2):
        for j, col in enumerate(cab, start=1):
            v = f.get(col, "")
            if col in NUMERICAS and str(v).strip():
                try:
                    v = int(float(v))
                except ValueError:
                    pass
            c = hoja.cell(row=i, column=j, value=v if v != "" else None)
            c._style = estilo[j]

    hoja.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{ultima}"
    hoja.freeze_panes = "D2"

    hoja.data_validations.dataValidation = []
    for coord, formula in validaciones:
        letra = re.match(r"([A-Z]+)", coord).group(1)
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        hoja.add_data_validation(dv)
        dv.add(f"{letra}2:{letra}{ultima}")

    # Las fórmulas del resumen fijan la última fila: se estiran.
    cambiadas = 0
    for h in wb.worksheets:
        if h is hoja:
            continue
        for fila in h.iter_rows():
            for c in fila:
                if isinstance(c.value, str) and "Catálogo!" in c.value:
                    nueva = re.sub(r"(Catálogo!\$[A-Z]+\$\d+:\$[A-Z]+\$)\d+",
                                   lambda m: m.group(1) + str(ultima), c.value)
                    if nueva != c.value:
                        c.value = nueva
                        cambiadas += 1
                elif isinstance(c.value, str) and f"las {previa - 1} series" in c.value:
                    c.value = c.value.replace(f"las {previa - 1} series", f"las {len(filas)} series")

    wb.save(XLSX)
    print(f"«Catálogo»: {previa - 1} -> {len(filas)} filas · {cambiadas} fórmulas estiradas")
    print(f"-> {XLSX.relative_to(RAIZ)}")


if __name__ == "__main__":
    main()
